from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook

from backend.app.learning.adaptive_engine.document_import import (
    AdaptiveDocumentImporter,
)

from backend.app.learning.adaptive_engine.question_bank import (
    StackXmlQuestionBankImporter,
)


def test_excel_question_bank_import() -> None:
    workbook = Workbook()

    sheet = workbook.active

    sheet.append(
        [
            "question_id",
            "question",
            "answer",
            "difficulty",
            "tags",
        ]
    )

    sheet.append(
        [
            "Q1",
            "Solve 2*x + 4 = 10",
            "3",
            -0.5,
            "algebra",
        ]
    )

    stream = BytesIO()

    workbook.save(
        stream
    )

    importer = (
        AdaptiveDocumentImporter()
    )

    converted = importer.import_bytes(
        filename="questions.xlsx",
        content=stream.getvalue(),
    )

    assert converted.question_count == 1

    bank = (
        StackXmlQuestionBankImporter()
        .import_text(
            converted.question_bank_xml
        )
    )

    assert bank.count == 1

    assert (
        bank.require(
            "Q1"
        )
        .adaptive_question
        .title
        == "Solve 2*x + 4 = 10"
    )


def test_word_question_bank_import() -> None:
    lines = [
        "Question ID: Q1",
        "Question: Solve x + 2 = 5",
        "Answer: 3",
        "Difficulty: -0.5",
        "Tags: algebra",
    ]

    paragraphs = "".join(
        (
            "<w:p>"
            "<w:r>"
            "<w:t>"
            + line
            + "</w:t>"
            "</w:r>"
            "</w:p>"
        )
        for line
        in lines
    )

    document_xml = (
        '<?xml version="1.0" '
        'encoding="UTF-8"?>'
        '<w:document '
        'xmlns:w="'
        'http://schemas.openxmlformats.org/'
        'wordprocessingml/2006/main'
        '">'
        "<w:body>"
        + paragraphs
        + "</w:body>"
        "</w:document>"
    )

    stream = BytesIO()

    with ZipFile(
        stream,
        "w",
        ZIP_DEFLATED,
    ) as archive:
        archive.writestr(
            "word/document.xml",
            document_xml,
        )

    converted = (
        AdaptiveDocumentImporter()
        .import_bytes(
            filename="questions.docx",
            content=stream.getvalue(),
        )
    )

    assert converted.question_count == 1

    assert (
        "Q1"
        in converted.question_bank_xml
    )


def test_metadata_is_generated() -> None:
    workbook = Workbook()

    sheet = workbook.active

    sheet.append(
        [
            "question_id",
            "question",
            "answer",
            "difficulty",
            "supports",
        ]
    )

    sheet.append(
        [
            "Q2",
            "Solve x - 4 = 3",
            "7",
            -0.2,
            "needs_linear_support",
        ]
    )

    stream = BytesIO()

    workbook.save(
        stream
    )

    converted = (
        AdaptiveDocumentImporter()
        .import_bytes(
            filename="questions.xlsx",
            content=stream.getvalue(),
        )
    )

    assert (
        "needs_linear_support"
        in converted.metadata_json
    )


def test_unsupported_file_rejected() -> None:
    importer = (
        AdaptiveDocumentImporter()
    )

    try:
        importer.import_bytes(
            filename="questions.txt",
            content=b"example",
        )
    except ValueError:
        return

    raise AssertionError(
        "Expected unsupported format "
        "to be rejected."
    )


def test_csv_question_bank_import() -> None:
    csv_text = (
        "question_id,question,answer,difficulty,tags\n"
        "Q10,Solve x + 1 = 4,3,-0.2,algebra\n"
    )

    converted = (
        AdaptiveDocumentImporter()
        .import_bytes(
            filename="questions.csv",
            content=csv_text.encode(
                "utf-8"
            ),
        )
    )

    assert converted.question_count == 1

    assert (
        "Q10"
        in converted.question_bank_xml
    )


def test_xlsm_question_bank_import() -> None:
    workbook = Workbook()

    sheet = workbook.active

    sheet.append(
        [
            "question_id",
            "question",
            "answer",
        ]
    )

    sheet.append(
        [
            "Q11",
            "Solve x + 5 = 9",
            "4",
        ]
    )

    stream = BytesIO()

    workbook.save(
        stream
    )

    converted = (
        AdaptiveDocumentImporter()
        .import_bytes(
            filename="questions.xlsm",
            content=stream.getvalue(),
        )
    )

    assert converted.question_count == 1

    assert (
        "Q11"
        in converted.question_bank_xml
    )
